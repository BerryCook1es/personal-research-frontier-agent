from __future__ import annotations

from pathlib import Path
from typing import Any


FIELDS = [
    "priority", "relevance_score", "research_track", "related_project", "potential_use",
    "recommended_action", "keyword_tier", "semantic_score", "publication_date", "venue",
    "title_cn", "title", "authors", "research_question", "method_summary", "main_contributions",
    "why_relevant", "abstract_cn", "abstract", "doi_url", "reading_status", "paper_id",
]


def render_excel(papers: list[dict[str, Any]], stats: dict[str, Any], output_path: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Excel output requires openpyxl") from exc
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Weekly Brief"
    summary.append(["Metric", "Count"])
    for key in ("scanned", "candidates", "A", "B", "C", "D", "Ignore"):
        summary.append([key, stats.get(key, 0)])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173A5E")

    def sheet(name: str, rows: list[dict[str, Any]]) -> None:
        ws = workbook.create_sheet(name)
        ws.append(FIELDS)
        for paper in rows:
            ws.append(["; ".join(map(str, paper.get(key, []))) if isinstance(paper.get(key), list) else paper.get(key, "") for key in FIELDS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173A5E")
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(55, max(12, max(len(str(c.value or "")) for c in column) + 2))
        doi_col = FIELDS.index("doi_url") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=doi_col)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    for priority in ("A", "B", "C", "D"):
        rows = [p for p in papers if p.get("priority") == priority]
        if rows:
            sheet(priority, rows)
    sheet("All", papers)
    workbook.save(output_path)
    return output_path

