#!/usr/bin/env python3
"""多格式输出渲染器：Excel / 交互HTML / Markdown笔记 / 终端预览。

Usage:
  python -X utf8 scripts/render_outputs.py --mode excel --input screened.json
  python -X utf8 scripts/render_outputs.py --mode app --input screened.json
  python -X utf8 scripts/render_outputs.py --mode notes --input screened.json
  python -X utf8 scripts/render_outputs.py --mode codex --input screened.json
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import sys
from datetime import date
from pathlib import Path

from _lib import POOL_LABELS, ROOT, TIER_LABELS, UI

OUTPUTS = ROOT / "outputs"
DISPLAY_MODES = ("excel", "app", "notes", "codex")
TIER_ORDER = {"core": 0, "proxy": 1, "eco": 2, "other": 3}


# ══════════════════════════════════════════════════════════════
# Excel 仪表盘
# ══════════════════════════════════════════════════════════════

def render_excel(papers: list[dict], stats: dict, output_path: Path | None = None) -> Path:
    """多 Sheet Excel 仪表盘，含条件格式和超链接。"""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    target = output_path or (OUTPUTS / "excel" / f"frontier_dashboard_{date.today().isoformat()}.xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # ── Sheet 1: Summary ──
    ws_summary = wb.active
    ws_summary.title = "Summary"
    summary_data = [
        ["统计项", "数量"],
        ["扫描论文总数", stats.get("total", 0)],
        ["核心 Core", stats.get("core", 0)],
        ["邻近 Proxy", stats.get("proxy", 0)],
        ["背景 Eco", stats.get("eco", 0)],
        ["其他 Other", stats.get("other", 0)],
    ]
    # 池统计
    pool_counts: dict[str, int] = {}
    for p in papers:
        pool = p.get("pool", "")
        if pool:
            pool_counts[pool] = pool_counts.get(pool, 0) + 1
    for pool, count in sorted(pool_counts.items(), key=lambda x: -x[1]):
        summary_data.append([f"  {POOL_LABELS.get(pool, pool)}", count])

    for row in summary_data:
        ws_summary.append(row)
    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 15
    for cell in ws_summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    # ── 通用表头 ──
    # header label → paper dict key mapping
    _header_map = [("tier", "tier"), ("pool", "pool"), ("publication_date", "publication_date"),
                   ("journal", "journal"), ("primary_topic", "primary_topic"),
                   ("all_topics", "all_topics"),
                   ("title_cn", "title_cn"), ("title_en", "title"),
                   ("authors", "authors"), ("abstract_cn", "abstract_cn"),
                   ("abstract_en", "abstract"), ("abstract_source", "abstract_source"),
                   ("doi_url", "doi_url")]
    headers = [h for h, _ in _header_map]

    def _write_sheet(ws, rows: list[dict], title: str):
        ws.title = title
        ws.append(headers)
        keys = [k for _, k in _header_map]
        for r in rows:
            vals = []
            for k in keys:
                v = r.get(k, "")
                vals.append("; ".join(v) if isinstance(v, list) else v)
            ws.append(vals)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        # 列宽
        widths = {"A": 8, "B": 14, "C": 14, "D": 28, "E": 22, "F": 30, "G": 50, "H": 50,
                  "I": 30, "J": 55, "K": 55, "L": 14, "M": 35}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        # DOI 超链接（M 列）
        for row in ws.iter_rows(min_row=2, max_col=len(headers)):
            cell = row[12]  # doi_url
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"
        # tier 条件格式
        max_r = max(2, ws.max_row)
        green = PatternFill("solid", fgColor="D9EAD3")
        yellow = PatternFill("solid", fgColor="FFF2CC")
        blue = PatternFill("solid", fgColor="CFE2F3")
        gray = PatternFill("solid", fgColor="E7E6E6")
        for col_letter, tier, fill in [("A", "core", green), ("A", "proxy", yellow),
                                        ("A", "eco", blue), ("A", "other", gray)]:
            ws.conditional_formatting.add(
                f"A2:A{max_r}",
                FormulaRule(formula=[f'$A2="{tier}"'], fill=fill),
            )

    # ── Sheet 2-5: 按 tier 分表 + 全量表 ──
    core_p = [p for p in papers if p.get("tier") == "core"]
    proxy_p = [p for p in papers if p.get("tier") == "proxy"]
    eco_p = [p for p in papers if p.get("tier") == "eco"]

    for rows, name in [(core_p, "Core"), (proxy_p, "Proxy"), (eco_p, "Eco")]:
        if rows:
            _write_sheet(wb.create_sheet(), sorted(rows, key=lambda r: r.get("publication_date", ""), reverse=True), name)

    # 全量表（按 tier + date 排序）
    all_sorted = sorted(papers, key=lambda r: (TIER_ORDER.get(r.get("tier", "other"), 9),
                                                r.get("publication_date", "")), reverse=False)
    if all_sorted:
        _write_sheet(wb.create_sheet(), all_sorted, "All")

    wb.save(target)
    return target


# ══════════════════════════════════════════════════════════════
# 交互 HTML 表格
# ══════════════════════════════════════════════════════════════

def render_app(papers: list[dict], stats: dict, output_path: Path | None = None,
              ref_date: str = "", from_date: str = "") -> Path:
    """生成交互式论文卡片界面（搜索 / 筛选 / 排序 / 摘要展开）。"""
    app_dir = output_path or (OUTPUTS / "app" / f"frontier_app_{date.today().isoformat()}")
    app_dir.mkdir(parents=True, exist_ok=True)

    # 精简数据，摘要保留完整内容供前端展开
    slim = []
    for p in papers:
        slim.append({
            "tier": p.get("tier", "other"),
            "tier_label": TIER_LABELS.get(p.get("tier", "other"), ""),
            "pool": POOL_LABELS.get(p.get("pool", ""), p.get("pool", "")),
            "pool_key": p.get("pool", ""),
            "date": p.get("publication_date", ""),
            "journal": p.get("journal", "") or p.get("journal_source", ""),
            "title_cn": p.get("title_cn", ""),
            "title_en": p.get("title", ""),
            "authors": p.get("authors", ""),
            "abstract_cn": (p.get("abstract_cn") or "").strip(),
            "abstract_en": (p.get("abstract") or "").strip(),
            "doi": p.get("doi_url", ""),
            "all_topics": p.get("all_topics", []),
            "all_topics_cn": p.get("all_topics_cn", []),
        })

    data_json = json.dumps(slim, ensure_ascii=False)

    total = stats.get("total", len(papers))
    core_n = stats.get("core", 0)
    proxy_n = stats.get("proxy", 0)
    eco_n = stats.get("eco", 0)
    other_n = stats.get("other", 0)

    index_html = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>前沿论文周报 · Frontier Tracker</title>
<style>
:root {{
  --c-core: #27ae60; --c-proxy: #f39c12; --c-eco: #2980b9; --c-other: #95a5a6;
  --bg: #f2f4f7; --card-bg: #ffffff; --text: #1e293b; --text2: #64748b;
  --border: #e2e8f0; --radius: 12px;
}}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ font-family: -apple-system, "Microsoft YaHei", "PingFang SC", "Noto Sans SC", system-ui, sans-serif; background: var(--bg); color: var(--text); line-height: 1.6; min-height: 100vh; -webkit-font-smoothing: antialiased; }}

/* ── Header ── */
.app-header {{ background: linear-gradient(135deg, #0f2b3d 0%, #1a4a6e 40%, #1e5a8a 100%); color: #fff; padding: 22px 32px; position: sticky; top: 0; z-index: 100; box-shadow: 0 2px 12px rgba(0,0,0,.15); }}
.app-header h1 {{ font-size: 1.35em; font-weight: 700; letter-spacing: -.2px; margin-bottom: 2px; }}
.app-header .subtitle {{ font-size: .78em; opacity: .65; font-weight: 400; }}
.app-header .stats-row {{ display: flex; gap: 10px; margin-top: 12px; flex-wrap: wrap; }}
.app-header .stat-chip {{ background: rgba(255,255,255,.12); padding: 4px 14px; border-radius: 20px; cursor: pointer; transition: all .2s; border: 1px solid transparent; font-size: .85em; font-weight: 500; }}
.app-header .stat-chip:hover {{ background: rgba(255,255,255,.22); transform: translateY(-1px); }}
.app-header .stat-chip.active {{ background: rgba(255,255,255,.28); border-color: rgba(255,255,255,.45); box-shadow: 0 0 0 2px rgba(255,255,255,.1); }}
.app-header .stat-chip .num {{ font-weight: 700; margin-left: 2px; opacity: .85; }}

/* ── Main ── */
main {{ max-width: 1100px; margin: 0 auto; padding: 20px 24px 48px; }}

/* ── Toolbar ── */
.toolbar {{ display: flex; gap: 10px; flex-wrap: wrap; align-items: center; margin-bottom: 14px; }}
.search-box {{ flex: 1; min-width: 260px; position: relative; }}
.search-box::before {{ content: '\\1F50D'; position: absolute; left: 13px; top: 50%; transform: translateY(-50%); font-size: 14px; z-index: 2; pointer-events: none; opacity: .5; }}
.search-box input {{ width: 100%; padding: 9px 14px 9px 38px; border: 1px solid var(--border); border-radius: 24px; font-size: 14px; background: var(--card-bg); transition: box-shadow .2s, border-color .2s; color: var(--text); }}
.search-box input::placeholder {{ color: #a0aec0; }}
.search-box input:focus {{ outline: none; border-color: #2980b9; box-shadow: 0 0 0 3px rgba(41,128,185,.10); }}
select {{ padding: 9px 30px 9px 12px; border: 1px solid var(--border); border-radius: 24px; font-size: 13px; background: var(--card-bg); cursor: pointer; appearance: none; background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' width='10' height='6'%3E%3Cpath d='M0 0l5 6 5-6z' fill='%2395a5a6'/%3E%3C/svg%3E"); background-repeat: no-repeat; background-position: right 12px center; color: var(--text); }}
select:focus {{ outline: none; border-color: #2980b9; }}

/* ── Control bar ── */
.bar {{ display: flex; gap: 8px; flex-wrap: wrap; align-items: center; margin-bottom: 16px; font-size: .87em; }}
.pill {{ padding: 6px 18px; border-radius: 20px; border: 1px solid var(--border); background: var(--card-bg); cursor: pointer; font-size: .87em; transition: all .18s; white-space: nowrap; color: var(--text2); font-weight: 500; }}
.pill:hover {{ border-color: #2980b9; color: #2980b9; }}
.pill.active {{ background: #1e5a8a; color: #fff; border-color: #1e5a8a; font-weight: 600; }}
.sort-btn {{ padding: 6px 14px; border-radius: 20px; border: 1px solid var(--border); background: var(--card-bg); cursor: pointer; font-size: .87em; transition: all .18s; white-space: nowrap; color: var(--text2); font-weight: 500; }}
.sort-btn:hover {{ border-color: #2980b9; color: #2980b9; }}
.sort-btn.active {{ background: #e8f0fe; border-color: #2980b9; color: #1a5276; font-weight: 600; }}
#count {{ color: var(--text2); font-size: .84em; font-weight: 500; }}
.spacer {{ flex: 1; }}

/* ── Paper cards ── */
.paper-list {{ display: flex; flex-direction: column; gap: 12px; }}
.paper-card {{ background: var(--card-bg); border-radius: var(--radius); box-shadow: 0 1px 4px rgba(0,0,0,.05); overflow: hidden; transition: box-shadow .25s, transform .15s; border-left: 5px solid transparent; position: relative; }}
.paper-card:hover {{ box-shadow: 0 4px 20px rgba(0,0,0,.10); transform: translateY(-2px); }}
.paper-card.tier-core {{ border-left-color: var(--c-core); }}
.paper-card.tier-proxy {{ border-left-color: var(--c-proxy); }}
.paper-card.tier-eco {{ border-left-color: var(--c-eco); }}
.paper-card.tier-other {{ border-left-color: var(--c-other); opacity: .82; }}
.card-body {{ padding: 16px 20px; }}
.card-badges {{ display: flex; gap: 8px; align-items: center; margin-bottom: 10px; flex-wrap: wrap; }}
.card-title-en {{ font-size: 1.06em; font-weight: 700; color: #1a3a4a; margin-bottom: 2px; line-height: 1.55; }}
.card-title-cn {{ font-size: .92em; color: #2c3e50; margin-bottom: 10px; line-height: 1.5; font-weight: 400; }}
.card-meta {{ display: flex; gap: 14px; flex-wrap: wrap; align-items: center; font-size: .8em; color: var(--text2); margin-bottom: 8px; }}
.card-meta-item {{ display: inline-flex; align-items: center; gap: 4px; }}
.card-meta-dot {{ width: 4px; height: 4px; border-radius: 50%; background: #cbd5e0; flex-shrink: 0; }}

/* ── Badges ── */
.badge {{ display: inline-block; padding: 3px 10px; border-radius: 12px; font-size: .75em; font-weight: 600; letter-spacing: .1px; }}
.badge-core {{ background: #d5f5e3; color: #1a6e34; }}
.badge-proxy {{ background: #fef3cd; color: #b8860b; }}
.badge-eco {{ background: #d6eaf8; color: #1a5276; }}
.badge-other {{ background: #e8e8e8; color: #777; }}
.badge-pool {{ background: #f1f5f9; color: #475569; font-weight: 500; font-size: .72em; }}
.badge-pool-patent-tech-mining {{ background: #fef3e3; color: #b8730e; }}
.badge-pool-bibliometrics-evaluation {{ background: #d6eaf8; color: #1a5276; }}
.badge-pool-science-policy-innovation {{ background: #d1f2eb; color: #0e6655; }}
.badge-pool-info-methods-datamining {{ background: #efe5f5; color: #6c3483; }}
.badge-pool-ai-usage-behavior {{ background: #fadbd8; color: #c0392b; }}
.badge-pool-complex-network-analysis {{ background: #d5f5e3; color: #1e8449; }}
.badge-pool-knowledge-graph-semantic {{ background: #e8eaed; color: #4a5568; }}
.badge-pool-comprehensive-high-impact {{ background: #fcf3cf; color: #7d6608; }}
.badge-topic {{ font-size: .70em; background: #e8f0fe; color: #1a5276; border: 1px solid #c5d9f0; }}
.badge-topic-cn {{ font-size: .68em; background: #fefce8; color: #7d6608; border: 1px solid #fde68a; }}
.topic-pair {{ display: inline-flex; gap: 0; align-items: stretch; margin-right: 5px; border-radius: 12px; overflow: hidden; border: 1px solid #dde; }}

/* ── Abstract ── */
.card-abstract {{ margin-top: 10px; padding-top: 10px; border-top: 1px solid #eef1f5; }}
.abs-cn {{ font-size: .88em; color: var(--text); line-height: 1.72; overflow: hidden; transition: max-height .35s ease; max-height: 4.6em; }}
.abs-cn.expanded {{ max-height: none; }}
.abs-en {{ font-size: .83em; color: var(--text2); line-height: 1.62; margin-top: 6px; overflow: hidden; transition: max-height .35s ease; max-height: 3.1em; }}
.abs-en.expanded {{ max-height: none; }}
.abs-toggle {{ display: inline-block; margin-top: 6px; font-size: .8em; color: #2980b9; cursor: pointer; user-select: none; font-weight: 600; transition: color .15s; }}
.abs-toggle:hover {{ color: #1a5276; }}
.no-abstract {{ color: #bcc4d0; font-size: .83em; }}

/* ── Card actions ── */
.card-actions {{ display: flex; gap: 16px; align-items: center; margin-top: 12px; }}
.card-actions a {{ font-size: .83em; color: #1e5a8a; text-decoration: none; font-weight: 600; transition: color .15s; }}
.card-actions a:hover {{ color: #2980b9; }}
.copy-doi {{ font-size: .83em; color: #94a3b8; cursor: pointer; border: none; background: none; padding: 2px 0; transition: color .15s; }}
.copy-doi:hover {{ color: #475569; }}

/* ── Toast ── */
.toast {{ position: fixed; bottom: 28px; left: 50%; transform: translateX(-50%); background: #1e293b; color: #fff; padding: 10px 24px; border-radius: 24px; font-size: .88em; z-index: 999; opacity: 0; transition: opacity .25s, transform .25s; pointer-events: none; box-shadow: 0 4px 16px rgba(0,0,0,.2); }}
.toast.show {{ opacity: 1; transform: translateX(-50%) translateY(-4px); }}

/* ── Empty state ── */
.empty {{ text-align: center; padding: 80px 20px; color: #bcc4d0; }}
.empty-icon {{ font-size: 3.5em; margin-bottom: 12px; }}

/* ── Responsive ── */
@media (max-width: 768px) {{
  .app-header {{ padding: 16px 18px; }}
  .app-header h1 {{ font-size: 1.15em; }}
  .app-header .stats-row {{ gap: 6px; }}
  .app-header .stat-chip {{ padding: 3px 10px; font-size: .78em; }}
  main {{ padding: 12px 10px 32px; }}
  .card-body {{ padding: 12px 14px; }}
  .card-title-cn {{ font-size: 1.0em; }}
  .toolbar {{ flex-direction: column; }}
  .search-box {{ min-width: 100%; }}
  select {{ width: 100%; }}
  .bar {{ flex-wrap: wrap; gap: 6px; }}
  .pill, .sort-btn {{ font-size: .82em; padding: 5px 12px; }}
}}

/* ── Scrollbar ── */
::-webkit-scrollbar {{ width: 6px; }}
::-webkit-scrollbar-thumb {{ background: #c4c4c4; border-radius: 3px; }}
::-webkit-scrollbar-track {{ background: transparent; }}
</style>
</head>
<body>
<div class="app-header">
  <h1>前沿论文周报 · Frontier Weekly</h1>
  <div class="subtitle">覆盖周期：{from_date or '?'} → {ref_date or date.today().isoformat()} · 专利计量 / 文献计量 / 技术挖掘 / 科学政策 / AI使用行为</div>
  <div class="stats-row">
    <span class="stat-chip active" data-tier="" onclick="filterByStatChip('',this)">全部<span class="num">{total}</span></span>
    <span class="stat-chip" data-tier="core" onclick="filterByStatChip('core',this)">核心<span class="num">{core_n}</span></span>
    <span class="stat-chip" data-tier="proxy" onclick="filterByStatChip('proxy',this)">邻近<span class="num">{proxy_n}</span></span>
    <span class="stat-chip" data-tier="eco" onclick="filterByStatChip('eco',this)">背景<span class="num">{eco_n}</span></span>
    <span class="stat-chip" data-tier="other" onclick="filterByStatChip('other',this)">其他<span class="num">{other_n}</span></span>
  </div>
</div>
<main>
  <div class="toolbar">
    <div class="search-box">
      <input id="q" placeholder="搜索标题、期刊、作者…（按 / 聚焦搜索）" oninput="render()">
    </div>
    <select id="pool" onchange="render()"><option value="">全部来源</option></select>
    <select id="journal" onchange="render()"><option value="">全部期刊</option></select>
  </div>
  <div class="bar">
    <button class="pill active" data-days="7" onclick="filterByDays(7,this)">近 7 天</button>
    <button class="pill" data-days="30" onclick="filterByDays(30,this)">近 30 天</button>
    <button class="pill" data-days="0" onclick="filterByDays(0,this)">全部时间</button>
    <span style="color:#cbd5e0;margin:0 4px;">|</span>
    <button class="sort-btn active" data-sort="date-desc" onclick="setSort('date-desc',this)">最新优先</button>
    <button class="sort-btn" data-sort="date-asc" onclick="setSort('date-asc',this)">最早优先</button>
    <button class="sort-btn" data-sort="tier" onclick="setSort('tier',this)">按级别</button>
    <span class="spacer"></span>
    <span id="count"></span>
  </div>
  <div id="list" class="paper-list"></div>
  <div id="toast" class="toast"></div>
</main>
<script>
const papers = {data_json};
const refDate = new Date('{ref_date or date.today().isoformat()}');
let filterDays = 7;
let filterTier = '';
let sortBy = 'date-desc';

document.addEventListener('DOMContentLoaded', function() {{
  initDropdowns();
  render();
  document.addEventListener('keydown', function(e) {{
    if (e.key === '/' && document.activeElement !== document.getElementById('q')) {{
      e.preventDefault();
      document.getElementById('q').focus();
    }}
  }});
}});

function initDropdowns() {{
  const pools = [...new Set(papers.map(p => p.pool).filter(Boolean))].sort();
  const journals = [...new Set(papers.map(p => p.journal).filter(Boolean))].sort();
  const poolEl = document.getElementById('pool');
  const journalEl = document.getElementById('journal');
  pools.forEach(v => {{ const o = document.createElement('option'); o.value = v; o.textContent = v; poolEl.appendChild(o); }});
  journals.forEach(v => {{ const o = document.createElement('option'); o.value = v; o.textContent = v; journalEl.appendChild(o); }});
}}

function filterByDays(days, btn) {{
  document.querySelectorAll('.pill').forEach(b => b.classList.remove('active'));
  btn.classList.add('active');
  filterDays = days;
  render();
}}

function filterByStatChip(tier, btn) {{
  document.querySelectorAll('.stat-chip').forEach(b => b.classList.remove('active'));
  btn.classList.add('active');
  filterTier = tier;
  render();
}}

function setSort(sort, btn) {{
  sortBy = sort;
  document.querySelectorAll('.sort-btn').forEach(b => b.classList.remove('active'));
  btn.classList.add('active');
  render();
}}

const TIER_ORDER = {{ core: 0, proxy: 1, eco: 2, other: 3 }};

function render() {{
  const q = document.getElementById('q').value.toLowerCase().trim();
  const pool = document.getElementById('pool').value;
  const journal = document.getElementById('journal').value;

  let filtered = papers.filter(p => {{
    if (filterDays > 0 && p.date) {{
      const diff = (refDate - new Date(p.date)) / 86400000;
      if (diff > filterDays) return false;
    }}
    if (filterTier && p.tier !== filterTier) return false;
    if (pool && p.pool !== pool) return false;
    if (journal && p.journal !== journal) return false;
    if (q) {{
      const hay = [p.title_cn, p.title_en, p.journal, p.authors, p.pool].join(' ').toLowerCase();
      if (!hay.includes(q)) return false;
    }}
    return true;
  }});

  filtered.sort((a, b) => {{
    if (sortBy === 'date-desc') return (b.date || '').localeCompare(a.date || '');
    if (sortBy === 'date-asc') return (a.date || '').localeCompare(b.date || '');
    if (sortBy === 'tier') return (TIER_ORDER[a.tier] || 9) - (TIER_ORDER[b.tier] || 9);
    return 0;
  }});

  document.getElementById('count').textContent = `${{filtered.length}} / ${{papers.length}} 篇`;

  const list = document.getElementById('list');
  if (filtered.length === 0) {{
    list.innerHTML = `<div class="empty"><div class="empty-icon">&#128236;</div><p>没有匹配的论文</p><p style="font-size:.82em;margin-top:4px;">试试调整筛选条件或搜索关键词</p></div>`;
    return;
  }}

  list.innerHTML = filtered.map((p, idx) => {{
    const hasAbs = p.abstract_cn || p.abstract_en;
    const shortAuthors = (p.authors || '').slice(0, 45);
    const authorsSuffix = (p.authors || '').length > 45 ? '…' : '';

    const topics = (p.all_topics || []).slice(0, 3);
    const topicsHtml = topics.map((t, i) => {{
      const cn = (p.all_topics_cn || [])[i] || '';
      if (cn) return `<span class="topic-pair"><span class="badge badge-topic">${{t}}</span><span class="badge badge-topic-cn">${{cn}}</span></span>`;
      return `<span class="badge badge-topic">${{t}}</span>`;
    }}).join('');

    const uid = 'p' + idx;

    let absHtml = '';
    if (hasAbs) {{
      absHtml = `<div class="card-abstract">
        ${{p.abstract_cn ? `<div class="abs-cn" id="${{uid}}-cn">${{p.abstract_cn}}</div>` : ''}}
        ${{p.abstract_en ? `<div class="abs-en" id="${{uid}}-en">${{p.abstract_en}}</div>` : ''}}
        <span class="abs-toggle" onclick="toggleAbs('${{uid}}')">展开全文 &#9662;</span>
      </div>`;
    }} else {{
      absHtml = `<div class="card-abstract"><div class="no-abstract">暂无摘要 / No abstract available</div></div>`;
    }}

    return `<div class="paper-card tier-${{p.tier}}" id="${{uid}}">
      <div class="card-body">
        <div class="card-badges">
          <span class="badge badge-${{p.tier}}">${{p.tier_label}}</span>
          ${{p.pool ? `<span class="badge badge-pool badge-pool-${{p.pool_key || ''}}">${{p.pool}}</span>` : ''}}
        </div>
        <div class="card-title-en">${{p.title_en || '(Untitled)'}}</div>
        ${{p.title_cn ? `<div class="card-title-cn">${{p.title_cn}}</div>` : ''}}
        <div class="card-meta">
          <span class="card-meta-item">${{p.date || '?'}}</span>
          <span class="card-meta-dot"></span>
          <span class="card-meta-item">${{p.journal || '?'}}</span>
          <span class="card-meta-dot"></span>
          <span class="card-meta-item" title="${{p.authors}}">${{shortAuthors}}${{authorsSuffix}}</span>
        </div>
        ${{topicsHtml ? `<div style="margin-bottom:2px">${{topicsHtml}}</div>` : ''}}
        ${{absHtml}}
        <div class="card-actions">
          ${{p.doi ? `<a href="${{p.doi}}" target="_blank" rel="noopener">查看原文 DOI</a>` : ''}}
          ${{p.doi ? `<button class="copy-doi" onclick="copyDoi('${{p.doi}}')">复制链接</button>` : ''}}
        </div>
      </div>
    </div>`;
  }}).join('');
}}

function toggleAbs(uid) {{
  const cn = document.getElementById(uid + '-cn');
  const en = document.getElementById(uid + '-en');
  const btn = document.querySelector(`#${{CSS.escape(uid)}} .abs-toggle`);
  if (!cn && !en) return;
  const isExpanded = (cn || en).classList.contains('expanded');
  if (cn) cn.classList.toggle('expanded', !isExpanded);
  if (en) en.classList.toggle('expanded', !isExpanded);
  btn.innerHTML = isExpanded ? '展开全文 &#9662;' : '收起 &#9650;';
}}

function showToast(msg) {{
  const t = document.getElementById('toast');
  t.textContent = msg;
  t.classList.add('show');
  clearTimeout(t._tid);
  t._tid = setTimeout(() => t.classList.remove('show'), 1800);
}}

function copyDoi(url) {{
  navigator.clipboard.writeText(url).then(() => {{
    showToast('DOI 链接已复制到剪贴板');
  }}).catch(() => {{
    const ta = document.createElement('textarea');
    ta.value = url; ta.style.position = 'fixed'; ta.style.opacity = 0;
    document.body.appendChild(ta); ta.select();
    document.execCommand('copy'); document.body.removeChild(ta);
    showToast('DOI 链接已复制到剪贴板');
  }});
}}
</script>
</body>
</html>"""
    (app_dir / "index.html").write_text(index_html, encoding="utf-8")
    return app_dir


# ══════════════════════════════════════════════════════════════
# Markdown 笔记（按 tier 分目录）
# ══════════════════════════════════════════════════════════════

def render_notes(papers: list[dict], stats: dict, output_path: Path | None = None) -> Path:
    """为每篇非 other 论文生成一个 .md 文件，按 tier 分目录。"""
    base = output_path or (OUTPUTS / "notes" / date.today().isoformat())
    base.mkdir(parents=True, exist_ok=True)

    written = 0
    for p in papers:
        tier = p.get("tier", "other")
        if tier == "other":
            continue

        tier_dir = base / tier
        tier_dir.mkdir(parents=True, exist_ok=True)

        title = p.get("title", "Untitled")
        slug = _slugify(title)
        name = f"{p.get('publication_date', 'nodate')}_{slug}.md"

        def _safe_yaml(v: str) -> str:
            """转义可能破坏 YAML frontmatter 的字符。"""
            return v.replace(chr(34), "'").replace("---", "\\u2014").replace("\n", " ")

        title_cn = p.get("title_cn", "")
        title_en = p.get("title", "")
        journal = p.get("journal", "") or p.get("journal_source", "")
        authors = p.get("authors", "")
        pub_date = p.get("publication_date", "")
        doi_url = p.get("doi_url", "")
        abstract_cn = (p.get("abstract_cn") or "").strip()
        abstract_en = (p.get("abstract") or "").strip()
        pool_label = POOL_LABELS.get(p.get("pool", ""), p.get("pool", ""))
        primary_topic = p.get("primary_topic", "")
        primary_topic_cn = p.get("primary_topic_cn", "")
        all_topics = p.get("all_topics", [])
        all_topics_cn = p.get("all_topics_cn", [])

        topics_str = ""
        if all_topics:
            pairs = list(zip(all_topics, all_topics_cn)) if all_topics_cn else [(t, "") for t in all_topics]
            topic_items = [f"- {en} / {cn}" if cn else f"- {en}" for en, cn in pairs[:5]]
            topics_str = "\n  ".join(topic_items)
        topic_frontmatter = ", ".join(all_topics[:3]) if all_topics else ""
        text = f"""---
title: "{_safe_yaml(title_en)}"
title_cn: "{_safe_yaml(title_cn)}"
journal: "{_safe_yaml(journal)}"
published: "{pub_date}"
doi: "{doi_url}"
tier: "{tier}"
pool: "{pool_label}"
authors: "{_safe_yaml(authors)}"
primary_topic: "{_safe_yaml(primary_topic)}"
primary_topic_cn: "{_safe_yaml(primary_topic_cn)}"
topics: "{_safe_yaml(topic_frontmatter)}"
---

# {title_cn or title_en}

## 基本信息
- **期刊**：{journal}
- **日期**：{pub_date}
- **DOI**：[{doi_url}]({doi_url})
- **作者**：{authors}
- **级别**：{TIER_LABELS.get(tier, tier)} | **来源**：{pool_label}
- **主主题**：{primary_topic} / {primary_topic_cn}
- **全部主题**：
  {topics_str or "(无)"}

## 英文标题
{title_en}

## 中文摘要
{abstract_cn or "(暂无)"}

## 英文摘要
{abstract_en or "(No abstract)"}

## 阅读笔记
- 与我研究的关系：
- 核心发现：
- 方法：
- 可引用点：
- 下一步：read / cite / note / archive
"""
        (tier_dir / name).write_text(text, encoding="utf-8")
        written += 1

    # 生成 CSV 索引
    csv_path = base / "_index.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        fieldnames = ["tier", "pool", "publication_date", "journal", "title", "title_cn", "authors", "doi_url"]
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for p in papers:
            if p.get("tier") == "other":
                continue
            writer.writerow({k: p.get(k, "") for k in fieldnames})

    print(f"  Notes: {written} papers → {base}", file=sys.stderr)
    return base


def _esc_md(v: str) -> str:
    """转义 markdown 表格中的 | 和换行。"""
    return v.replace("|", "\\|").replace("\n", " ")


def _slugify(text: str, max_len: int = 80) -> str:
    """中文/英文标题 → 文件名安全 slug。"""
    import re
    value = re.sub(r"[^a-zA-Z0-9一-鿿]+", "-", text).strip("-")
    return value[:max_len] or "paper"


# ══════════════════════════════════════════════════════════════
# 终端 / 文件预览
# ══════════════════════════════════════════════════════════════

def render_codex(papers: list[dict], stats: dict, output_path: Path | None = None,
                 limit: int = 50, to_stdout: bool = True) -> Path:
    """生成 Markdown 预览，按 tier 分组。"""
    target = output_path or (OUTPUTS / "codex" / f"frontier_preview_{date.today().isoformat()}.md")
    target.parent.mkdir(parents=True, exist_ok=True)

    lines = [
        f"# {UI['title']} Preview",
        f"Generated: {date.today().isoformat()}",
        f"Total: {stats.get('total', '?')} | Core: {stats.get('core', '?')} "
        f"| Proxy: {stats.get('proxy', '?')} | Eco: {stats.get('eco', '?')}",
        "",
    ]

    for tier, heading in [(t, f"## {TIER_LABELS[t]}") for t in ("core", "proxy", "eco")]:
        tier_papers = [p for p in papers if p.get("tier") == tier]
        if not tier_papers:
            continue
        lines.append(heading)
        lines.append("| Date | Journal | Title (CN) | Title (EN) | DOI |")
        lines.append("| --- | --- | --- | --- | --- |")
        for p in tier_papers[:limit]:
            date_s = p.get("publication_date", "")
            journal = p.get("journal", "") or p.get("journal_source", "")
            title_cn = (p.get("title_cn") or p.get("title", ""))[:60]
            title_en = (p.get("title") or "")[:80]
            doi = p.get("doi_url", "")
            lines.append(
                f"| {_esc_md(date_s)} | {_esc_md(journal)} | "
                f"{_esc_md(title_cn)} | {_esc_md(title_en)} | {_esc_md(doi)} |"
            )
        lines.append("")

    content = "\n".join(lines) + "\n"
    target.write_text(content, encoding="utf-8")

    if to_stdout:
        print(content)

    return target


# ══════════════════════════════════════════════════════════════
# CLI
# ══════════════════════════════════════════════════════════════

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Render screened data into display formats.")
    parser.add_argument("--mode", choices=DISPLAY_MODES, required=True, help="Output format.")
    parser.add_argument("--input", type=Path, required=True, help="Screened JSON file.")
    parser.add_argument("--output", type=Path, help="Output file or directory path.")
    parser.add_argument("--limit", type=int, default=50, help="Max papers per tier in codex mode.")
    return parser


def main() -> int:
    args = build_parser().parse_args()

    if not args.input.exists():
        print(f"[ERROR] Input file not found: {args.input}", file=sys.stderr)
        return 1

    data = json.loads(args.input.read_text(encoding="utf-8"))
    papers = data.get("new", []) + data.get("already_seen", [])
    stats = data.get("_profile_stats", {})

    if args.mode == "excel":
        result = render_excel(papers, stats, args.output)
    elif args.mode == "app":
        result = render_app(papers, stats, args.output)
    elif args.mode == "notes":
        result = render_notes(papers, stats, args.output)
    else:
        result = render_codex(papers, stats, args.output, args.limit)

    print(json.dumps({
        "mode": args.mode, "papers": len(papers),
        "output": str(result),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

